from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

doc = Document()
section = doc.sections[0]
section.top_margin = section.bottom_margin = Inches(0.75)
styles = doc.styles
styles['Normal'].font.name = 'Aptos'
styles['Normal'].font.size = Pt(11)
title = doc.add_paragraph(style='Title')
title.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = title.add_run('Northwind Services Employee Guide')
run.font.color.rgb = None
doc.add_paragraph('This short guide explains the company policies used for the document question answering demo. It covers leave, expenses, remote work, and support contacts.')

for heading, body in [
    ('Annual Leave', 'Full-time employees receive 24 days of paid annual leave per calendar year. Submit leave requests at least 10 business days before the first day of leave. Your manager must approve the request before travel is booked.'),
    ('Business Expenses', 'Submit business expenses within 21 days of purchase through the expense portal. Attach an itemized receipt for any expense above $30. Meals are reimbursable up to $45 per person when they relate to approved business travel.'),
    ('Remote Work', 'Employees may work remotely up to three days each week with manager approval. Team members must be available between 10:00 and 16:00 local time on working days.'),
    ('Support', 'For payroll questions, contact payroll@northwind.example. For IT access problems, submit a ticket through the internal help desk.'),
]:
    doc.add_heading(heading, level=1)
    doc.add_paragraph(body)

doc.core_properties.title = 'Northwind Services Employee Guide'
doc.core_properties.author = 'Northwind Services'
doc.save('sample_employee_guide.docx')
